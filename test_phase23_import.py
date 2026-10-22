"""المرحلة 23 — استيراد الأطباء والمعامل من Excel. قاعدة مؤقتة فقط."""
import os, sys, json, sqlite3, threading, urllib.request, urllib.error, tempfile, time, base64, io, zipfile
_T = tempfile.mkdtemp(prefix="lq23-"); os.environ["LAQINO_DB"] = _T + "/t.db"; os.environ["LAQINO_UPLOADS"] = _T + "/up"; os.makedirs(_T + "/up")
if os.environ["LAQINO_DB"].startswith(os.getcwd() + "/laqino.db"): print("ISOLATION BROKEN"); sys.exit(1)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import server, testutil, bulk_import
from http.server import ThreadingHTTPServer
import openpyxl
ok = lambda n, c: print(("PASS " if c else "FAIL ") + n) or (c or sys.exit(1))
server.init("01000000000", "AdminPass123")
srv = ThreadingHTTPServer(("127.0.0.1", 0), server.H); threading.Thread(target=srv.serve_forever, daemon=True).start(); time.sleep(.3)
BASE = "http://127.0.0.1:%d/api" % srv.server_address[1]
def call(m, p, d=None, t=None):
    r = urllib.request.Request(BASE + p, json.dumps(d).encode() if d is not None else None, {"Content-Type": "application/json", **({"Authorization": "Bearer " + t} if t else {})}, method=m)
    try: x = urllib.request.urlopen(r); return x.status, json.loads(x.read())
    except urllib.error.HTTPError as e: return e.code, json.loads(e.read())
at = call("POST", "/login", {"phone": "01000000000", "password": "AdminPass123"})[1]["token"]
call("POST", "/register", {"phone": "01022222222", "name": "عميل", "password": "password1"})
ct = call("POST", "/login", {"phone": "01022222222", "password": "password1"})[1]["token"]
mk = lambda n, ar, act=1: call("POST", "/admin/regions", {"parent_id": ar, "level": "village", "slug": n, "name_ar": n_ar[n], "is_active": bool(act)}, at)[1]["id"]
n_ar = {"a": "قرية الأمل", "dup": "النخيل", "dup2": "النخيل", "off": "قرية موقوفة"}
rA = mk("a", 3); rD1 = call("POST", "/admin/regions", {"parent_id": 3, "level": "village", "slug": "d1", "name_ar": "النخيل", "is_active": True}, at)[1]["id"]
c0 = call("POST", "/admin/regions", {"parent_id": 2, "level": "markaz", "slug": "mk2", "name_ar": "مركز سيدي سالم", "is_active": True}, at)[1]["id"]
rD2 = call("POST", "/admin/regions", {"parent_id": c0, "level": "village", "slug": "d2", "name_ar": "النخيل", "is_active": True}, at)[1]["id"]
rOff = call("POST", "/admin/regions", {"parent_id": 3, "level": "village", "slug": "off", "name_ar": "قرية موقوفة", "is_active": False}, at)[1]["id"]

def book(docs, labs, hdr_row=1, names=("الأطباء", "المعامل")):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = names[0]
    for _ in range(hdr_row - 1): ws.append(["جدول أطباء المنطقة"])
    for r in docs:
        if r: ws.append(r)
        else: ws._current_row += 1   # صف فارغ حقيقي (Excel لا يكتبه)
    if labs is not None:
        w2 = wb.create_sheet(names[1])
        for r in labs: w2.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()
b64 = lambda raw: base64.b64encode(raw).decode()
DH = ["الاسم", "التخصص", "المنطقة", "المركز", "العيادة", "العنوان", "المواعيد", "الهاتف", "واتساب", "نبذة"]
LH = ["اسم المعمل", "المنطقة", "العنوان", "المواعيد", "التحاليل", "سحب منزلي", "الهاتف"]
docs = [DH,
  ["د. أحمد محمود", "باطنة", "قرية الأمل", "", "عيادة النور", "شارع المحطة", "السبت 5م", 1012345678, None, ""],        # رقم بلا صفر أول (Excel)
  ["د. سارة علي", "أطفال", "قرية الأمل", "", "", "", "", "٠١١١٢٢٢٣٣٣٣", "", ""],                                       # أرقام عربية، ناقص عنوان
  ["د. محمد حسن", "جلدية", "النخيل", "مركز سيدي سالم", "", "الشارع الرئيسي", "", "01201112223 / 01001112223", "", ""],  # مكرر الاسم → المركز يحسمها + رقمان
  ["د. منى سمير", "جلدية", "النخيل", "", "", "عنوان", "", "01012345679", "", ""],                                       # غامضة بلا مركز → خطأ
  ["د. هاني فؤاد", "عظام", "قرية موقوفة", "", "", "", "", "01012345670", "", ""],                                        # منطقة غير مفعّلة → خطأ
  ["د. كريم", "", "قرية الأمل", "", "", "", "", "01012345671", "", ""],                                                  # بلا تخصص → خطأ
  ["د. ليلى نبيل", "أنف وأذن", "قرية الأمل", "", "", "ع", "", "12345", "", ""],                                          # رقم غير صالح → خطأ
  ["د. أحمد محمود", "باطنة", "قرية الأمل", "", "", "", "", "", "", ""],                                                   # مكرر داخل الملف
  [], ["د. رامي", "باطنة", "غير موجودة", "", "", "", "", "", "", ""],]
labs = [LH, ["معمل الأمل", "قرية الأمل", "بجوار الوحدة", "يوميًا", "سكر، كبد", "نعم", "01012345680"], ["معمل بلا عنوان", "قرية الأمل", "", "", "", "لا", "01012345681"]]
f = book(docs, labs, hdr_row=3)   # صف العناوين في الصف 3
snap = lambda: sqlite3.connect(os.environ["LAQINO_DB"]).execute("SELECT (SELECT COUNT(*) FROM doctors),(SELECT COUNT(*) FROM labs),(SELECT COUNT(*) FROM specialties)").fetchone()
s0 = snap()

# ---- صلاحيات
ok("template/preview/commit require login", all(call(m, p, {"data": ""})[0] == 401 for m, p in (("GET", "/admin/import/template"), ("POST", "/admin/import/preview"), ("POST", "/admin/import/commit"))))
ok("non-admin customer gets 403", call("GET", "/admin/import/template", None, ct)[0] == 403 and call("POST", "/admin/import/commit", {"data": b64(f)}, ct)[0] == 403)

# ---- القالب يعمل ذهابًا وإيابًا
st, t = call("GET", "/admin/import/template", None, at)
wbt = openpyxl.load_workbook(io.BytesIO(base64.b64decode(t["data"])))
ok("template is a valid xlsx with the two sheets and RTL", st == 200 and wbt.sheetnames == ["الأطباء", "المعامل"] and wbt["الأطباء"].sheet_view.rightToLeft)
st, pv = call("POST", "/admin/import/preview", {"data": t["data"], "default_region_id": rA}, at)
ok("template parses (both sheets seen); its placeholder region is reported, not silently accepted", st == 200 and pv["summary"]["doctors"]["error"] == 1 and pv["summary"]["labs"]["error"] == 1 and "المنطقة غير موجودة" in pv["items"][0]["msg"])

# ---- المعاينة لا تكتب شيئًا
st, pv = call("POST", "/admin/import/preview", {"data": b64(f)}, at)
sm = pv["summary"]
ok("preview: 200 and finds both sheets despite header on row 3", st == 200 and sm["doctors"]["new"] + sm["doctors"]["error"] + sm["doctors"]["duplicate"] == 9 and sm["labs"]["new"] == 2)
ok("preview writes nothing", snap() == s0)
by = {i["name"] + "|" + str(i["row"]): i for i in pv["items"]}
find = lambda nm: [i for i in pv["items"] if i["name"] == nm]
ok("valid doctors are new (3: digits-less phone, Arabic digits, markaz-disambiguated)", sm["doctors"]["new"] == 3)
ok("ambiguous village without markaz → error asking for markaz/default", "مكرر" in find("د. منى سمير")[0]["msg"])
ok("inactive region → error", "غير مفعّلة" in find("د. هاني فؤاد")[0]["msg"])
ok("missing specialty → error", find("د. كريم")[0]["msg"] == "التخصص مطلوب")
ok("bad phone → error", "رقم غير صالح" in find("د. ليلى نبيل")[0]["msg"])
ok("unknown region → error", "غير موجودة" in find("د. رامي")[0]["msg"])
ok("same doctor twice in the file → second is duplicate", [i["status"] for i in find("د. أحمد محمود")] == ["new", "duplicate"] and "داخل الملف" in find("د. أحمد محمود")[1]["msg"])
ok("two numbers in one cell → first kept with a warning", any("أول" in w for w in find("د. محمد حسن")[0]["warn"]))
ok("new specialties are announced (باطنة، أطفال، جلدية)", set(pv["new_specialties"]) == {"باطنة", "أطفال", "جلدية"})
ok("blank rows ignored", not [i for i in pv["items"] if not i["name"] and i["status"] != "error"] and all(i["row"] != 12 for i in pv["items"]) and any(i["row"] == 13 for i in pv["items"]))

# ---- التنفيذ
ok("commit of a non-xlsx file is rejected cleanly", call("POST", "/admin/import/commit", {"data": b64(b"hello")}, at)[0] == 400 and snap() == s0)
ok("garbage base64 rejected", call("POST", "/admin/import/commit", {"data": "@@@"}, at)[0] == 400)
ok("bad default region rejected", call("POST", "/admin/import/preview", {"data": b64(f), "default_region_id": rOff}, at)[0] == 400)
st, cm = call("POST", "/admin/import/commit", {"data": b64(f)}, at)
ok("commit ok: 3 doctors + 2 labs + 3 specialties, nothing published", st == 200 and cm["done"]["doctors"] == 3 and cm["done"]["labs"] == 2 and cm["done"]["specialties"] == 3 and cm["done"]["published"] == 0)
db = sqlite3.connect(os.environ["LAQINO_DB"]); db.row_factory = sqlite3.Row
D = {r["name_ar"]: r for r in db.execute("SELECT d.*,s.name_ar sp FROM doctors d JOIN specialties s ON s.id=d.specialty_id")}
ok("specialties created and linked automatically", D["د. أحمد محمود"]["sp"] == "باطنة" and D["د. سارة علي"]["sp"] == "أطفال")
ok("phone fixed: leading zero restored + Arabic digits converted", D["د. أحمد محمود"]["phone"] == "01012345678" and D["د. سارة علي"]["phone"] == "01112223333")
ok("markaz column picked the right النخيل", D["د. محمد حسن"]["region_id"] == rD2 and D["د. أحمد محمود"]["region_id"] == rA)
ok("everything lands unpublished", all(r["is_published"] == 0 for r in D.values()) and db.execute("SELECT COUNT(*) FROM labs WHERE is_published=1").fetchone()[0] == 0)
L = {r["name_ar"]: r for r in db.execute("SELECT * FROM labs")}
ok("lab fields + home visit parsed", L["معمل الأمل"]["home_visit"] == 1 and L["معمل الأمل"]["tests"] == "سكر، كبد" and L["معمل بلا عنوان"]["home_visit"] == 0)
ok("audit entry written", db.execute("SELECT COUNT(*) FROM audit_log WHERE action='import.excel'").fetchone()[0] == 1)
ok("imported rows are visible to the existing admin lists", len(call("GET", "/admin/doctors", None, at)[1]) == 3 and len(call("GET", "/admin/labs", None, at)[1]) == 2)
ok("specialty list now contains the new ones", {"باطنة", "أطفال", "جلدية"} <= {x["name_ar"] for x in call("GET", "/specialties?all=1", None, at)[1]})

# ---- إعادة الرفع: لا تكرار
st, again = call("POST", "/admin/import/commit", {"data": b64(f)}, at)
ok("re-importing the same file adds nothing (409)", st == 409 and snap()[0] == 3 and snap()[1] == 2)
st, pv2 = call("POST", "/admin/import/preview", {"data": b64(f)}, at)
ok("re-preview shows existing as duplicates", pv2["summary"]["doctors"]["new"] == 0 and pv2["summary"]["doctors"]["duplicate"] >= 4 and pv2["summary"]["labs"]["duplicate"] == 2)

# ---- تحديث الموجود يلمس الخلايا غير الفارغة فقط + النشر
call("POST", "/admin/doctors/%d" % D["د. أحمد محمود"]["id"], {"bio": "نبذة يدوية"}, at)
f2 = book([DH, ["د. أحمد محمود", "باطنة", "قرية الأمل", "", "", "", "الأحد 6م", "", "", ""],
               ["د. وليد صبري", "قلب", "قرية الأمل", "", "عيادة القلب", "ميدان", "يوميًا", "01012345690", "", ""],
               ["د. نادية", "باطنة", "قرية الأمل", "", "", "", "", "", "", ""]], [LH])
st, c2 = call("POST", "/admin/import/commit", {"data": b64(f2), "update_existing": True, "publish": True}, at)
db = sqlite3.connect(os.environ["LAQINO_DB"]); db.row_factory = sqlite3.Row
a = db.execute("SELECT * FROM doctors WHERE name_ar='د. أحمد محمود'").fetchone()
ok("update_existing: schedule updated, blank cells did NOT erase phone/clinic/bio", st == 200 and a["schedule"] == "الأحد 6م" and a["phone"] == "01012345678" and a["clinic_name"] == "عيادة النور" and a["bio"] == "نبذة يدوية")
w = db.execute("SELECT * FROM doctors WHERE name_ar='د. وليد صبري'").fetchone(); n = db.execute("SELECT * FROM doctors WHERE name_ar='د. نادية'").fetchone()
ok("publish: complete new row published, incomplete one stays hidden, existing visibility untouched", w["is_published"] == 1 and n["is_published"] == 0 and a["is_published"] == 0 and c2["done"]["published"] == 1)
ok("published doctor appears in the public list; unpublished ones do not", [x["name_ar"] for x in call("GET", "/doctors?specialty_id=%d" % w["specialty_id"])[1]] == ["د. وليد صبري"])

# ---- معاملة واحدة: لو فشل التنفيذ لا يبقى شيء
before = snap(); orig = bulk_import.audit
bulk_import.audit = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom"))
f3 = book([DH, ["د. اختبار الفشل", "تخصص جديد تمامًا", "قرية الأمل", "", "", "ع", "", "01012345699", "", ""]], None)
try: st = call("POST", "/admin/import/commit", {"data": b64(f3)}, at)[0]
except Exception: st = 500
bulk_import.audit = orig
ok("failure mid-commit rolls back everything (no doctor, no specialty left behind)", snap() == before)

# ---- اختيار الصفحات بالمحتوى، ومنطقة افتراضية، وملف ناقص
f4 = book([["الاسم", "القسم", "الموبايل"], ["د. بدون منطقة", "أسنان", "01012345611"]], None, names=("ورقة1", "x"))
st, p4 = call("POST", "/admin/import/preview", {"data": b64(f4)}, at)
ok("sheet detected by columns (specialty column → doctors); missing region reported", st == 200 and p4["summary"]["doctors"]["error"] == 1 and p4["items"][0]["msg"] == "المنطقة مطلوبة")
st, p4 = call("POST", "/admin/import/preview", {"data": b64(f4), "default_region_id": rA}, at)
ok("default region fills blank region cells (synonym headers القسم/الموبايل)", p4["summary"]["doctors"]["new"] == 1)
f5 = book([["لا شيء", "هنا"], ["a", "b"]], None, names=("ورقة", "y"))
ok("workbook with no recognisable sheet → 400 with guidance", call("POST", "/admin/import/preview", {"data": b64(f5)}, at)[0] == 400)
bomb = io.BytesIO()
with zipfile.ZipFile(bomb, "w") as z: z.writestr("xl/workbook.xml", '<!DOCTYPE x [<!ENTITY a "b">]><workbook/>')
ok("XML with DOCTYPE/ENTITY rejected", call("POST", "/admin/import/preview", {"data": b64(bomb.getvalue())}, at)[0] == 400)

# ---- لا تأثر بالمسارات القائمة
st, d1 = call("POST", "/admin/doctors", {"name_ar": "د. يدوي", "specialty_id": w["specialty_id"], "region_id": rA, "address": "x", "phone": "01012345600"}, at)
ok("manual add + publish still works exactly as before", st == 201 and call("POST", "/admin/doctors/%d/publish" % d1["id"], {"published": True}, at)[0] == 200)
ok("manual lab add still works", call("POST", "/admin/labs", {"name_ar": "معمل يدوي", "region_id": rA, "address": "x"}, at)[0] == 201)
print("ALL PASS")
