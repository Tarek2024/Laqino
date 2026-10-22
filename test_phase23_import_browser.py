"""المرحلة 23 — واجهة استيراد Excel في لوحة الإدارة (Chromium حقيقي، قاعدة مؤقتة، هاتف 390px)."""
import os, sys, json, time, tempfile, threading, io, urllib.request, urllib.error
try:
    from playwright.sync_api import sync_playwright
    import openpyxl
except Exception:
    print("SKIP browser test (playwright/openpyxl not installed)"); sys.exit(0)
_T = tempfile.mkdtemp(prefix="lq23b-"); os.environ["LAQINO_DB"] = _T + "/t.db"; os.environ["LAQINO_UPLOADS"] = _T + "/up"; os.makedirs(_T + "/up")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import server
from http.server import ThreadingHTTPServer
server.init("01000000000", "AdminPass123")
srv = ThreadingHTTPServer(("127.0.0.1", 0), server.H); threading.Thread(target=srv.serve_forever, daemon=True).start(); time.sleep(.3)
BASE = "http://127.0.0.1:%d" % srv.server_address[1]
ok = lambda n, c: print(("PASS " if c else "FAIL ") + n) or (c or sys.exit(1))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "الأطباء"
ws.append(["الاسم", "التخصص", "المنطقة", "العنوان", "الهاتف"]); ws.append(["د. أحمد محمود", "باطنة", "مركز الرياض", "شارع المحطة", 1012345678]); ws.append(["د. بلا تخصص", "", "مركز الرياض", "x", "01012345600"])
w2 = wb.create_sheet("المعامل"); w2.append(["الاسم", "المنطقة", "العنوان", "سحب منزلي"]); w2.append(["معمل الأمل", "مركز الرياض", "بجوار الوحدة", "نعم"])
f = _T + "/in.xlsx"; wb.save(f)
with sync_playwright() as p:
    b = p.chromium.launch(); pg = b.new_context(viewport={"width": 390, "height": 844}, locale="ar-EG", service_workers="block", accept_downloads=True).new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(BASE + "/admin.html"); pg.fill("#p", "01000000000"); pg.fill("#w", "AdminPass123"); pg.click("#app button"); pg.wait_for_selector("nav")
    for tab, label in (("doctors", "الأطباء"), ("labs", "المعامل")):
        pg.click("nav button[data-k=%s]" % tab); pg.wait_for_selector("#impcard", timeout=8000)
        ok("[%s] import card appended under the existing screen" % tab, pg.locator("#impcard").count() == 1 and pg.locator("#v .card").count() >= 2)
        ok("[%s] original add form still there" % tab, pg.locator("#v input[placeholder^='اسم']").count() >= 1)
    with pg.expect_download() as d: pg.click("text=تنزيل نموذج Excel")
    d.value.save_as(_T + "/tpl.xlsx")
    ok("template downloads as .xlsx", d.value.suggested_filename.endswith(".xlsx") and openpyxl.load_workbook(_T + "/tpl.xlsx").sheetnames == ["الأطباء", "المعامل"])
    pg.set_input_files("#imf", f); pg.wait_for_selector("#imres table", timeout=8000)
    t = pg.inner_text("#imres")
    ok("preview shows counts, the new specialty and the row error", "الأطباء: جديد 1" in t and "أخطاء 1" in t and "باطنة" in t and "التخصص مطلوب" in t and "المعامل: جديد 1" in t)
    ok("preview wrote nothing", json.loads(urllib.request.urlopen(urllib.request.Request(BASE + "/api/admin/doctors", headers={"Authorization": "Bearer " + pg.evaluate("sessionStorage.getItem('at')")})).read()) == [])
    ok("no sideways page scroll at 390px", pg.evaluate("document.documentElement.scrollWidth<=innerWidth+1"))
    pg.click("#impgo"); pg.wait_for_selector("text=تم الاستيراد", timeout=8000)
    ok("commit result shown (unpublished by default)", "أُضيف 1 طبيب و1 معمل" in pg.inner_text("#imres") and "غير منشورة" in pg.inner_text("#imres"))
    tok = pg.evaluate("sessionStorage.getItem('at')"); H = {"Authorization": "Bearer " + tok}
    docs = json.loads(urllib.request.urlopen(urllib.request.Request(BASE + "/api/admin/doctors", headers=H)).read())
    ok("doctor saved via the UI with fixed phone, unpublished", len(docs) == 1 and docs[0]["phone"] == "01012345678" and docs[0]["is_published"] == 0 and docs[0]["specialty"] == "باطنة")
    ok("no JS errors", not errs)
    b.close()
print("ALL PASS")
